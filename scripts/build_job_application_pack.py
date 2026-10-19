from __future__ import annotations

import csv
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "job-applications" / "2026-08-18"
DOCX_DIR = OUT / "docx"
PDF_DIR = OUT / "pdf"
QA_DIR = OUT / "qa"
for folder in (DOCX_DIR, PDF_DIR, QA_DIR):
    folder.mkdir(parents=True, exist_ok=True)

NAVY = RGBColor(11, 37, 69)
BLUE = RGBColor(30, 93, 163)
INK = RGBColor(28, 34, 42)
MUTED = RGBColor(87, 97, 110)
LIGHT = "E8EEF5"

CONTACT = {
    "name": "Hamdi Jawher",
    "title": "Senior Product Designer | UX/UI | Design Systems",
    "location": "Tunis, Tunisia | Open to international relocation and EMEA remote work",
    "phone": "+216 22 085 367",
    "email": "hamdijawher@icloud.com",
    "portfolio": "https://www.jawherhamdi.com/work.html?filter=product",
    "linkedin": "https://www.linkedin.com/in/jawher-hamdi-748349189/",
}

ROLES = [
    {
        "slug": "xpin-ai-product-designer-dubai",
        "company": "Xpin AI",
        "role": "Product Designer - UI/UX",
        "location": "Dubai, UAE",
        "url": "https://www.linkedin.com/jobs/view/4446657655/",
        "source": "LinkedIn Easy Apply",
        "match": 90,
        "mobility": "Explicit UAE visa sponsorship and relocation support",
        "risk": "Portfolio craft is the primary filter; the product case studies must carry the application.",
        "focus": "sovereign AI, bilingual Arabic/English experiences, information-dense interfaces, and system-level RTL design",
        "hook": "Your brief is unusually close to work I have already done: designing multilingual, accessibility-led interfaces for diverse populations, including RTL behavior, while making complex and consequential systems understandable.",
        "evidence": "At TLScontact, I led UX/UI for biometric self-service experiences, researched users across 10 international visa centres, and designed multilingual patterns spanning six or more languages. At FORVIA, I simplified data-heavy Palantir Foundry workflows and helped establish a reusable design system. That combination maps directly to Xpin AI's need for dense enterprise interfaces, bilingual foundations, clear states, and close design-engineering collaboration.",
    },
    {
        "slug": "dust-senior-product-design-paris",
        "company": "Dust",
        "role": "Senior Product Designer",
        "location": "Paris, France",
        "url": "https://jobs.ashbyhq.com/dust/a58a9807-026c-469d-a550-582565b54329",
        "source": "Ashby",
        "match": 88,
        "mobility": "Explicit visa assistance and up to EUR 10,000 in Paris relocation support",
        "risk": "Elite AI startup with a very high craft bar and an office-first culture four days per week.",
        "focus": "defining AI-agent interaction patterns, complex B2B discovery, polished UI, technical collaboration, and end-to-end product ownership",
        "hook": "Dust's brief fits the strongest intersection in my background: simplify technically complex B2B workflows, work closely with engineers, and bring enough visual craft to define a product category whose interaction patterns are still emerging.",
        "evidence": "At FORVIA, I spent more than two years auditing and redesigning data-heavy Palantir Foundry applications, building reusable patterns, and collaborating with technical teams on complex operational flows. At TLScontact, I own international research, accessibility, multilingual behavior, high-fidelity prototyping, and executive alignment for a biometric kiosk. Through Jay Studio, I add art direction and brand judgment, giving me the generalist range Dust asks for across discovery, product, visual design, prototyping, and light technical work.",
    },
    {
        "slug": "superchat-ai-senior-product-designer-dubai",
        "company": "Superchat AI",
        "role": "Senior Product Designer",
        "location": "Dubai, UAE",
        "url": "https://www.linkedin.com/jobs/view/4454558714/",
        "source": "LinkedIn Easy Apply",
        "match": 87,
        "mobility": "Dubai on-site; sponsorship is not stated in the posting",
        "risk": "AI-product experience is preferred, so the application must lead with complex data products and practical AI-assisted workflows without overstating shipped AI features.",
        "focus": "AI-first SaaS experiences, complex workflows, research, analytics, scalable design systems, accessibility, and close product-engineering collaboration",
        "hook": "Superchat's scope maps directly to the work I have done across complex enterprise interfaces: own the problem end to end, translate dense technical requirements into clear flows, validate with users and data, and keep the design system strong enough to scale delivery.",
        "evidence": "At FORVIA, I redesigned Palantir Foundry data applications and developed reusable system patterns for complex dashboards. At Addinn, I used research, usability testing, product analytics, and A/B evidence across fintech and SaaS products. At TLScontact, I lead accessibility-led, multilingual service design and work across research, prototyping, delivery, and stakeholder alignment. I also bring HTML/CSS literacy and AI-assisted design workflows while staying precise about what has and has not shipped.",
    },
    {
        "slug": "salt-product-uiux-designer-abu-dhabi",
        "company": "Salt (client company)",
        "role": "Product & UI/UX Designer",
        "location": "Abu Dhabi, UAE",
        "url": "https://www.linkedin.com/jobs/view/4450875490/",
        "source": "LinkedIn Easy Apply",
        "match": 94,
        "mobility": "Abu Dhabi on-site five days per week; sponsorship is not stated",
        "risk": "The client is undisclosed and the role requires some 3D-tool exposure that is not explicit in the resume.",
        "focus": "enterprise product UX, high-stakes digital identity and security interfaces, design systems, executive presentations, websites, brand, and creative direction",
        "hook": "This role is the cleanest product-and-creative hybrid in the search: it needs someone who can simplify serious enterprise technology while also owning visual identity, websites, campaigns, presentations, and the quality of execution with developers.",
        "evidence": "My product experience covers data-heavy enterprise applications, fintech, biometric services, research, accessibility, design systems, and executive stakeholder communication. Through Jay Studio, I also develop visual systems, treatments, photography, film, campaign assets, and creative direction. That combination lets me move from user journey and interaction model to a polished interface, an executive presentation, or a public-facing brand touchpoint without losing coherence or implementation discipline.",
    },
    {
        "slug": "mistral-ai-product-designer-paris",
        "company": "Mistral AI",
        "role": "Product Designer",
        "location": "Paris, France / London, UK",
        "url": "https://jobs.ashbyhq.com/mistral.ai/f7e940ab-b60d-4ac2-9c34-aca245b8c750",
        "source": "Ashby",
        "match": 82,
        "mobility": "Role lists relocation support among location-dependent benefits",
        "risk": "Highly competitive AI employer; the portfolio needs clearer shipped outcomes and AI-product examples.",
        "focus": "making advanced AI useful, simple, and trustworthy across enterprise and self-serve product experiences",
        "hook": "Mistral's emphasis on useful, simple, and trustworthy AI reflects the standard I apply to any complex system: expose the right information, remove avoidable friction, and make the user's next action unmistakable.",
        "evidence": "My background spans enterprise data products, regulated financial flows, and international public-facing services. At FORVIA, I redesigned Palantir Foundry applications and built design-system guidance for complex dashboards. At TLScontact, I research users across multiple countries and turn operational constraints into accessible, multilingual prototypes. I can own research, interaction design, high-fidelity UI, systems documentation, and stakeholder communication from problem framing to delivery.",
    },
    {
        "slug": "geotab-senior-product-designer-dubai",
        "company": "Geotab",
        "role": "Senior Product Designer",
        "location": "Dubai, UAE",
        "url": "https://job-boards.greenhouse.io/geotab/jobs/5371579008",
        "source": "Greenhouse",
        "match": 87,
        "mobility": "Dubai-based role; sponsorship not stated in the posting",
        "risk": "Over 100 applicants and no explicit relocation statement.",
        "focus": "AI video intelligence, fleet dashboards, complex data-heavy surfaces, and automotive product design",
        "hook": "Geotab is one of the clearest matches in this search because my experience combines automotive data applications, complex dashboards, research, and design-system work.",
        "evidence": "For more than two years at FORVIA, I improved Palantir Foundry data applications, translating complex operational workflows into clearer dashboards and reusable patterns. I then moved into biometric service design at TLScontact, where I lead research, accessibility, prototyping, and cross-country stakeholder alignment. I can bring the autonomy, visual judgment, and systems thinking needed for live video review, safety timelines, and fleet intelligence surfaces.",
    },
    {
        "slug": "hostaway-staff-product-designer-design-systems-emea",
        "company": "Hostaway",
        "role": "Staff Product Designer - Design Systems",
        "location": "Remote EMEA",
        "url": "https://www.linkedin.com/jobs/view/4448742553/",
        "source": "LinkedIn / company application",
        "match": 80,
        "mobility": "Explicitly open to candidates physically based anywhere in EMEA",
        "risk": "Staff scope is a stretch; the resume lacks explicit people-leadership and production token-pipeline ownership.",
        "focus": "owning a design-system roadmap, reusable components, documentation, adoption, and the bridge between design and engineering",
        "hook": "Hostaway's design-system mandate matches the strongest systems work in my background: turning fragmented interfaces into documented, reusable patterns that help designers and engineers move faster without lowering quality.",
        "evidence": "At FORVIA, I developed and implemented a design system for Palantir Foundry applications, documenting patterns, usage guidance, and governance for data-heavy interfaces. At TLScontact, I apply structured critique, accessibility requirements, multilingual behavior, and usability validation to a complex kiosk ecosystem. I understand that adoption is the real measure of a design system, and I am comfortable pairing system work with hands-on product delivery.",
    },
    {
        "slug": "investengine-senior-product-designer-growth-emea",
        "company": "InvestEngine",
        "role": "Senior Product Designer - Growth",
        "location": "Remote EMEA (contract)",
        "url": "https://careers.investengine.com/jobs/8209727-senior-product-designer-growth",
        "source": "Company careers",
        "match": 88,
        "mobility": "Remote EMEA contract",
        "risk": "Contract role; compensation and contracting entity need confirmation.",
        "focus": "build-ready fintech journeys, regulated clarity, experimentation, behavioral evidence, and low-churn delivery",
        "hook": "The expectation to deliver high-risk customer journeys with strong upfront clarity and low rework maps directly to my fintech and evidence-led product background.",
        "evidence": "At Addinn, I led discovery and design for Smartweb Pay, combining interviews, UX audits, journey mapping, prototyping, usability testing, closed-beta analytics, and A/B results. At FORVIA, I partnered with technical teams on data-heavy applications and reusable systems. I work comfortably across product, engineering, research, edge cases, and implementation detail, with a strong bias toward customer comprehension rather than decorative polish.",
    },
    {
        "slug": "httpie-senior-product-designer-devtools-prague",
        "company": "HTTPie",
        "role": "Senior Product Designer - DevTools",
        "location": "Prague / Remote",
        "url": "https://jobs.ashbyhq.com/httpie/c98ab093-1fab-4f21-b4f5-b2af9071c1d0",
        "source": "Ashby",
        "match": 79,
        "mobility": "Remote role associated with Prague; employment geography must be confirmed",
        "risk": "Direct developer-tooling and API-domain depth are limited in the portfolio.",
        "focus": "simplifying constrained technical workflows, validating interaction models, and owning product touchpoints in a remote startup",
        "hook": "HTTPie's focus on simplicity in developer tooling is compelling because the core design problem is familiar to me: preserve expert capability while removing needless cognitive load.",
        "evidence": "At FORVIA, I worked inside a deeply technical Palantir Foundry environment, auditing complex flows, designing data applications, and creating reusable patterns in close collaboration with developers. My current TLScontact work requires translating operational and technical constraints into clear interaction models and tested prototypes. I also bring HTML/CSS literacy, structured written communication, and the curiosity needed to learn a technical domain quickly.",
    },
    {
        "slug": "palantir-product-designer-london",
        "company": "Palantir Technologies",
        "role": "Product Designer",
        "location": "London, UK",
        "url": "https://jobs.lever.co/palantir/2d8b7eaa-5c08-44ef-a0f0-5a96cee2794d",
        "source": "Lever",
        "match": 92,
        "mobility": "Hybrid London; sponsorship not stated",
        "risk": "The job stresses developer workflows and an exceptionally selective hiring bar.",
        "focus": "designing analytical software, micro-interactions, reusable components, research, and close partnership with engineers",
        "hook": "I am applying with unusually direct context: I spent more than two years designing and improving data applications built on Palantir Foundry within FORVIA.",
        "evidence": "That work included end-to-end UX audits, complex dashboard and workflow design, stakeholder discovery, iterative validation, and a design system for designers and developers. I learned how much clarity, consistency, and interaction detail matter when users make decisions through dense operational data. My current TLScontact role adds international research, accessibility, multilingual systems, and executive alignment. Together, those experiences prepare me to contribute across interaction design, prototyping, research, and engineering partnership at Palantir.",
    },
    {
        "slug": "shory-senior-ux-designer-abu-dhabi",
        "company": "Shory",
        "role": "Senior UX Designer",
        "location": "Abu Dhabi, UAE",
        "url": "https://www.linkedin.com/jobs/view/4454238111/",
        "source": "LinkedIn Easy Apply",
        "match": 92,
        "mobility": "Abu Dhabi on-site; sponsorship is not stated in the posting",
        "risk": "The posting is highly competitive and the application must make relocation and sponsorship needs explicit.",
        "focus": "simplifying digital insurance journeys through research, analytics, accessible design systems, AI-assisted prototyping, and Arabic/English RTL experience",
        "hook": "Shory's mission is unusually close to my background: turn complex financial-service journeys into clear, trustworthy digital experiences while working across research, interaction design, analytics, accessibility, and delivery.",
        "evidence": "At Addinn, I led discovery and design for Smartweb Pay, combining interviews, UX audits, journey mapping, prototyping, usability testing, closed-beta analytics, and A/B evidence. At TLScontact, I design multilingual, accessibility-led self-service experiences across international visa centres, including RTL behavior. At FORVIA, I simplified complex B2B data applications and built reusable design-system patterns. I also bring HTML/CSS literacy and practical AI-assisted workflows, matching the role's end-to-end and technically collaborative scope.",
    },
    {
        "slug": "easygenerator-product-designer-dubai",
        "company": "Easygenerator",
        "role": "Product Designer",
        "location": "Dubai, UAE (hybrid)",
        "url": "https://jobs.ashbyhq.com/Easygenerator/7123fb73-2e61-483f-b79b-4ebf0ec9497b",
        "source": "Ashby",
        "match": 90,
        "mobility": "Dubai hybrid; relocation and sponsorship are not stated in the posting",
        "risk": "Application requires a truthful example of a product decision that was wrong and how it was corrected; this cannot be invented from the resume.",
        "focus": "end-to-end B2B SaaS design, complex editors and dashboards, research, analytics, design systems, product copy, and AI-assisted implementation",
        "hook": "Easygenerator's scope maps tightly to how I work: own discovery through shipped quality, partner directly with product and engineering, use research and analytics, and improve a complex B2B product without designing around the system.",
        "evidence": "At FORVIA, I audited and redesigned complex Palantir Foundry dashboards and established reusable design-system patterns for data-heavy workflows. At Addinn, I combined research, prototyping, usability testing, product analytics, and A/B evidence across fintech and SaaS products. My current TLScontact work adds accessibility, multilingual behavior, international stakeholder alignment, and high-fidelity delivery. I write interface rationale, understand HTML/CSS, and use AI-assisted workflows while keeping product judgment and implementation quality explicit.",
    },
    {
        "slug": "way-senior-product-designer-paris",
        "company": "Way",
        "role": "Senior Product Designer",
        "location": "Paris, France",
        "url": "https://jobs.lever.co/Way/8f11267a-0be5-4e05-8a4f-5372bb582f0e",
        "source": "Lever",
        "match": 83,
        "mobility": "Paris role; sponsorship not stated",
        "risk": "Hospitality platform experience is indirect and relocation terms are not published.",
        "focus": "B2B experience technology for hospitality brands, combining product rigor with strong visual and experiential judgment",
        "hook": "Way sits at a useful intersection in my background: rigorous B2B product design on one side, and hands-on creative direction for hospitality and culture-led experiences on the other.",
        "evidence": "My product work spans research, complex workflows, dashboards, financial services, design systems, accessibility, and international stakeholder alignment. Through Jay Studio, I also develop visual concepts, treatments, film, and photography for hospitality and Tunisian brands. That combination helps me think beyond interface completion toward the full guest, operator, and brand experience while still delivering practical, build-ready product work.",
    },
    {
        "slug": "alan-product-designer-paris",
        "company": "Alan",
        "role": "Product Designer",
        "location": "Paris / selected European hubs",
        "url": "https://jobs.ashbyhq.com/alan/1bf22834-f85d-4581-95a8-0a0278ba9013",
        "source": "Ashby",
        "match": 82,
        "mobility": "Must be within commuting distance of a listed hub; relocation support not stated",
        "risk": "Hub-location requirement may block the application without sponsorship.",
        "focus": "member-centered health experiences, end-to-end product ownership, continuous testing, and design-system contribution",
        "hook": "Alan's ambition to make care more accessible and preventive resonates with my current work designing a public-facing service where confidence, accessibility, and cultural context are part of the product, not afterthoughts.",
        "evidence": "At TLScontact, I lead design for biometric self-service experiences, apply WCAG 2.1 AA, support multilingual and RTL behavior, and research users across 10 visa centres. Earlier roles added financial products, enterprise data applications, analytics, testing, and design systems. I bring full-stack product design skills across research, interaction, visual design, prototyping, delivery, and quality monitoring, along with the communication needed to align international teams.",
    },
    {
        "slug": "aspexx-senior-product-designer-riyadh",
        "company": "aspexx. (client consultancy)",
        "role": "Senior Product Designer",
        "location": "Riyadh, Saudi Arabia",
        "url": "https://www.linkedin.com/jobs/view/4450201739/",
        "source": "LinkedIn Easy Apply",
        "match": 86,
        "mobility": "Recruiter advertises visa, annual flights, private medical, and 25 days leave",
        "risk": "Client identity is undisclosed; contract terms and exact portfolio expectations need recruiter confirmation.",
        "focus": "consulting across government, banking, retail, hospitality, and automotive with strong end-to-end UX/UI delivery",
        "hook": "The client mix behind this role closely matches my experience across automotive data platforms, financial products, public-facing services, and hospitality-led creative work.",
        "evidence": "I bring more than five years of end-to-end product design across discovery, workshops, research, user journeys, prototyping, high-fidelity UI, testing, analytics, and design systems. FORVIA gave me deep experience with automotive and Palantir Foundry data applications; Addinn covered fintech and SaaS; TLScontact adds international service design and executive alignment. I am comfortable entering complex client environments quickly and translating business, regulatory, cultural, and technical constraints into coherent experiences.",
    },
]


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tag = OxmlElement("w:tblHeader")
    tag.set(qn("w:val"), "true")
    tr_pr.append(tag)


def set_table_widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)


def configure_document(doc: Document, title_text: str, subtitle_text: str | None = None):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.3)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, color, before, after in (
        ("Heading 1", 13.5, BLUE, 10, 4),
        ("Heading 2", 11.5, NAVY, 7, 2),
        ("Heading 3", 10.3, NAVY, 5, 1),
    ):
        st = styles[name]
        st.font.name = "Arial"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.text = "HAMDI JAWHER  |  SENIOR PRODUCT DESIGNER"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)
    for run in header.runs:
        run.font.name = "Arial"
        run.font.size = Pt(8)
        run.font.color.rgb = MUTED

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_after = Pt(0)
    run = footer.add_run("jawherhamdi.com  |  hamdijawher@icloud.com")
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_before = Pt(3)
    title.paragraph_format.space_after = Pt(2)
    run = title.add_run(title_text)
    run.font.name = "Arial"
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = NAVY
    if subtitle_text:
        sub = doc.add_paragraph()
        sub.paragraph_format.space_after = Pt(10)
        rr = sub.add_run(subtitle_text)
        rr.font.name = "Arial"
        rr.font.size = Pt(10.5)
        rr.font.color.rgb = MUTED


def add_contact_line(doc: Document):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    text = f"{CONTACT['phone']}  |  {CONTACT['email']}  |  jawherhamdi.com  |  LinkedIn: jawher-hamdi-748349189"
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(9)
    r.font.color.rgb = MUTED


def add_bullet(doc: Document, text: str, compact: bool = True):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.14)
    p.paragraph_format.space_after = Pt(2.2 if compact else 4)
    p.paragraph_format.line_spacing = 1.03
    p.add_run(text)


def build_resume():
    doc = Document()
    configure_document(doc, CONTACT["name"], CONTACT["title"])
    add_contact_line(doc)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.add_run(
        "Senior product and UX designer with 5+ years delivering complex enterprise, data-heavy, financial, and service experiences. "
        "Leads work from discovery and research through interaction design, high-fidelity prototyping, design systems, usability validation, and delivery. "
        "Combines rigorous product thinking with strong visual judgment, accessibility practice, and international stakeholder communication."
    )

    doc.add_heading("Core strengths", level=1)
    skills = doc.add_table(rows=2, cols=2)
    skills.style = "Table Grid"
    skills.autofit = False
    set_table_widths(skills, [3.42, 3.42])
    items = [
        ("Product design", "Discovery, journeys, interaction design, responsive UI, prototypes, edge cases, handoff"),
        ("Research", "Interviews, workshops, qualitative and quantitative research, usability testing, A/B testing"),
        ("Systems", "Figma libraries, reusable components, patterns, documentation, governance, accessibility"),
        ("Domains", "Enterprise data, Palantir Foundry, fintech, biometrics, automotive, multilingual services"),
    ]
    for cell, (label, detail) in zip([c for row in skills.rows for c in row.cells], items):
        cell.text = ""
        pp = cell.paragraphs[0]
        pp.paragraph_format.space_after = Pt(1)
        rr = pp.add_run(label + "\n")
        rr.bold = True
        rr.font.color.rgb = NAVY
        rr2 = pp.add_run(detail)
        rr2.font.size = Pt(9)

    doc.add_heading("Experience", level=1)
    experiences = [
        (
            "Senior UX Designer | TLScontact",
            "Aug 2025 - Present | Tunisia / international programmes",
            [
                "Lead end-to-end UX/UI for a self-service biometric kiosk, improving clarity, accessibility, and transaction flow across 5+ visa centres.",
                "Conduct research across 10 international visa centres and translate cultural, behavioural, and operational findings into interaction models and prototypes.",
                "Apply WCAG 2.1 AA principles and design multilingual patterns supporting 6+ languages, including RTL behavior.",
                "Present evidence-based recommendations to executives and business owners across multiple countries, aligning product, operations, and delivery teams.",
            ],
        ),
        (
            "UX/UI Designer | FORVIA",
            "May 2023 - Aug 2025 | Palantir Foundry data experiences",
            [
                "Directed UX audits and experience improvements for data applications built with Palantir Foundry, simplifying complex operational workflows and dashboards.",
                "Developed and implemented reusable design-system patterns, documentation, usage guidance, and governance for designers and developers.",
                "Designed data-exploration and decision-support experiences through stakeholder discovery, iterative prototyping, and user feedback.",
                "Facilitated an R&D ideation workshop in Paris-Nanterre, guiding teams from problem framing to actionable experience concepts.",
            ],
        ),
        (
            "UX/UI Product Designer | Addinn Group",
            "Oct 2020 - May 2023 | Fintech, HR SaaS, and service products",
            [
                "Led discovery and product design for Smartweb Pay, SynapseHR, and Lait's Collecte, working directly with product owners and stakeholders.",
                "Used interviews, UX audits, journey mapping, prototyping, user testing, closed-beta analytics, and A/B results to improve product decisions.",
                "Designed secure, comprehensible financial flows that balanced customer confidence, business goals, and technical constraints.",
            ],
        ),
        (
            "Independent Art Director & Creative Producer | Jay Studio",
            "Independent practice | Tunisia / international availability",
            [
                "Develop creative concepts, visual systems, treatments, storyboards, production plans, film, photography, and rollout assets for hospitality and culture-led brands.",
            ],
        ),
    ]
    for experience_index, (title, meta, bullets) in enumerate(experiences):
        if experience_index == 2:
            doc.add_page_break()
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(5)
        h.paragraph_format.space_after = Pt(0)
        h.paragraph_format.keep_with_next = True
        rr = h.add_run(title)
        rr.bold = True
        rr.font.size = Pt(10.5)
        rr.font.color.rgb = NAVY
        mm = doc.add_paragraph(meta)
        mm.paragraph_format.space_after = Pt(1)
        mm.paragraph_format.keep_with_next = True
        mm.runs[0].font.size = Pt(8.8)
        mm.runs[0].font.color.rgb = MUTED
        for bullet in bullets:
            add_bullet(doc, bullet)

    doc.add_heading("Selected product case studies", level=1)
    for case in (
        "Biometric Kiosk - international service research, multilingual and RTL behavior, WCAG-led interaction design, and high-fidelity prototyping.",
        "FORVIA Palantir Apps - UX audits, complex data workflows, dashboard interaction design, and reusable design-system patterns.",
        "Smart Web Pay - secure fintech journeys shaped through research, usability testing, closed-beta analytics, and A/B evidence.",
    ):
        add_bullet(doc, case, compact=False)

    doc.add_heading("Education, credentials, tools, and languages", level=1)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.add_run("Education: ").bold = True
    p.add_run("Bachelor Degree in Design, Higher Institute of Arts and Multimedia Manouba (ISAMM)")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.add_run("Credentials: ").bold = True
    p.add_run("Google UX Design Certificate; Meta Principles of UX/UI Design; IE Branding and Customer Experience; UXCEL Verified Designer (96/100, top 5%)")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.add_run("Tools: ").bold = True
    p.add_run("Figma, Adobe Creative Suite, Miro, Palantir Foundry, HTML/CSS, AI-assisted workflows, DaVinci Resolve, Final Cut Pro")
    p = doc.add_paragraph()
    p.add_run("Languages: ").bold = True
    p.add_run("English C2 | French B2")

    path = DOCX_DIR / "Hamdi_Jawher_Senior_Product_Designer_CV.docx"
    doc.save(path)
    return path


def build_cover_letter(role):
    doc = Document()
    configure_document(doc, "Application", f"{role['role']} | {role['company']} | 18 August 2026")
    add_contact_line(doc)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    p.add_run("Dear Hiring Team,")

    paragraphs = [
        f"I am applying for the {role['role']} position at {role['company']}. {role['hook']}",
        role["evidence"],
        f"What draws me to this role is the opportunity to focus on {role['focus']}. My portfolio shows both product case studies and the visual craft I bring from independent art-direction work. I would welcome the opportunity to discuss how I can contribute from the first discovery conversation through shipped, measurable product quality.",
        f"I am currently based in Tunis and am open to relocation or the stated remote arrangement. Portfolio: {CONTACT['portfolio']} | LinkedIn: {CONTACT['linkedin']}",
    ]
    for text in paragraphs:
        p = doc.add_paragraph(text)
        p.paragraph_format.space_after = Pt(9)
        p.paragraph_format.line_spacing = 1.12

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.add_run("Sincerely,\nHamdi Jawher\n").bold = True
    p.add_run(f"{CONTACT['email']} | {CONTACT['phone']}")

    path = DOCX_DIR / f"Cover_Letter_{role['slug']}.docx"
    doc.save(path)
    return path


def write_reports():
    md = [
        "# International job application report",
        "",
        "Prepared: 18 August 2026",
        "",
        "Important: the percentages below are requirement-match scores, not hiring probabilities. No honest recruiter can promise a 70-75% offer chance from a cold application.",
        "",
        "## Candidate positioning",
        "",
        "Primary lane: senior product/UX design for complex enterprise, data-heavy, regulated, and international services. Secondary differentiator: art direction, visual craft, film/production, and hospitality storytelling.",
        "",
        "## Application targets",
        "",
        "| # | Company | Role | Location | Match | Mobility signal | Status |",
        "|---:|---|---|---|---:|---|---|",
    ]
    for i, role in enumerate(ROLES, 1):
        md.append(
            f"| {i} | [{role['company']}]({role['url']}) | {role['role']} | {role['location']} | {role['match']}% | {role['mobility']} | Prepared - submission confirmation pending |"
        )
    md += ["", "## Role-by-role rationale and risk", ""]
    for i, role in enumerate(ROLES, 1):
        md += [
            f"### {i}. {role['company']} - {role['role']}",
            "",
            f"- Source: {role['source']}",
            f"- Vacancy: {role['url']}",
            f"- Match rationale: {role['hook']}",
            f"- Mobility: {role['mobility']}",
            f"- Hard risk: {role['risk']}",
            f"- Prepared file: docx/Cover_Letter_{role['slug']}.docx and pdf/Cover_Letter_{role['slug']}.pdf",
            "",
        ]
    md += [
        "## Naming and consistency issue",
        "",
        "The resume uses Hamdi Jawher, LinkedIn uses Jawher Hamdi, and the portfolio brand uses Jay Hamdi. Formal application files use Hamdi Jawher to stay consistent with the provided resume and email. This mismatch should be resolved before final interviews and legal onboarding.",
        "",
        "## Submission policy status",
        "",
        "The applications are prepared but not submitted in this report version. Resume upload, contact-data entry, job-form submission, recruiter messages, and emails transmit personal data and represent the candidate. They require a final action-time confirmation naming the destinations.",
    ]
    (OUT / "APPLICATION_REPORT.md").write_text("\n".join(md), encoding="utf-8")

    email_md = ["# Tailored recruiter email drafts", ""]
    for i, role in enumerate(ROLES, 1):
        email_md += [
            f"## {i}. {role['company']} - {role['role']}",
            "",
            f"**Subject:** Application - {role['role']} - Hamdi Jawher",
            "",
            "Dear Hiring Team,",
            "",
            role["hook"],
            "",
            f"{role['evidence']} I am currently based in Tunis and am open to the advertised relocation or remote arrangement.",
            "",
            f"Portfolio: {CONTACT['portfolio']}  ",
            f"LinkedIn: {CONTACT['linkedin']}  ",
            f"Email: {CONTACT['email']} | Phone: {CONTACT['phone']}",
            "",
            "Best,  ",
            "Hamdi Jawher",
            "",
        ]
    (OUT / "RECRUITER_EMAIL_DRAFTS.md").write_text("\n".join(email_md), encoding="utf-8")

    with (OUT / "application_tracker.csv").open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["company", "role", "location", "match_score", "mobility", "risk", "source", "url", "status", "cover_letter_pdf"])
        for role in ROLES:
            writer.writerow([
                role["company"], role["role"], role["location"], role["match"], role["mobility"], role["risk"], role["source"], role["url"],
                "Prepared - submission confirmation pending", f"pdf/Cover_Letter_{role['slug']}.pdf"
            ])


def main():
    build_resume()
    for role in ROLES:
        build_cover_letter(role)
    write_reports()
    print(f"Created {len(ROLES) + 1} DOCX files in {DOCX_DIR}")


if __name__ == "__main__":
    main()
